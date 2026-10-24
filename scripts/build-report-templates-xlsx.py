#!/usr/bin/env python3
"""Editable Excel versions of the eight free NDT report templates (owner, 2026-10-24).

Owner: "Report should be written as Atlantis Engineering Consultants with no logo and in excel
sheet, for them to download templates" — so visitors download directly instead of asking one by one.

Source of truth for the fields and example values is the same pair of files the template pages
already publish: public/templates/print/<slug>.html (sections, field labels, table columns) and
public/templates/csv/<slug>.csv (example values). Output: public/templates/xlsx/<slug>.xlsx.

Run: python3 scripts/build-report-templates-xlsx.py   (then scripts/recalc via the xlsx skill)
"""
import csv, html, re, sys
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
PRINT = ROOT / "public/templates/print"
CSVD = ROOT / "public/templates/csv"
OUT = ROOT / "public/templates/xlsx"
COMPANY = "Atlantis Engineering Consultants"
SITE = "atlantisndt.com"
QUOTE_URL = "https://atlantisndt.com/contact?service=templates&subject=Custom%20NDT%20report%20templates"

SLUGS = [
    "ut-weld-inspection-report", "ut-thickness-survey-report", "pt-liquid-penetrant-report",
    "mt-magnetic-particle-report", "rt-radiography-report", "vt-visual-inspection-report",
    "paut-phased-array-report", "tofd-report",
]

F = "Arial"
thin = Side(style="thin", color="555555")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
LABEL_FILL = PatternFill("solid", fgColor="EEF2F7")
INPUT_FILL = PatternFill("solid", fgColor="FFF9DB")   # light yellow = cells to fill in
HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
EX_FONT = Font(name=F, size=9, italic=True, color="555555")
WRAP = Alignment(wrap_text=True, vertical="top")
NCOLS_MIN = 8


def parse_print(slug):
    s = (PRINT / f"{slug}.html").read_text(encoding="utf-8")
    title = html.unescape(re.search(r"<h1>(.*?)</h1>", s).group(1))
    sections = []
    for m in re.finditer(r'<h2>(.*?)</h2><table class="(fields|grid)">(.*?)</table>', s, re.S):
        name, kind, body = html.unescape(m.group(1)), m.group(2), m.group(3)
        if kind == "grid":
            cols = [html.unescape(re.sub("<[^>]+>", "", c)) for c in re.findall(r"<th[^>]*>(.*?)</th>", body)]
            sections.append(("grid", name, cols))
        else:
            labels = [html.unescape(re.sub("<[^>]+>", "", c)) for c in re.findall(r"<th[^>]*>(.*?)</th>", body)]
            sections.append(("fields" if labels else "box", name, labels))
    return title, sections


def parse_examples(slug):
    rows = list(csv.reader((CSVD / f"{slug}.csv").open(encoding="utf-8-sig")))
    fields, grid = {}, []
    in_grid = False
    for r in rows[1:]:
        if not r or not any(c.strip() for c in r):
            in_grid = False
            continue
        if len(r) == 1:              # "Indication record (example values)" banner
            in_grid = True
            continue
        if in_grid:
            grid.append(r)
            continue
        if len(r) >= 3:
            fields[r[1]] = r[2]
    return fields, grid[1:] if grid else []   # first grid row is the header


def style_cell(c, *, bold=False, size=10, fill=None, color="111111", align=WRAP, border=True):
    c.font = Font(name=F, size=size, bold=bold, color=color)
    c.alignment = align
    if fill:
        c.fill = fill
    if border:
        c.border = BORDER


def build(slug):
    title, sections = parse_print(slug)
    ex_fields, ex_grid = parse_examples(slug)
    grid_cols = next((cols for kind, _, cols in sections if kind == "grid"), [])
    ncols = max(NCOLS_MIN, len(grid_cols) + (2 if slug == "ut-thickness-survey-report" else 0))
    last = get_column_letter(ncols)

    wb = Workbook()
    ws = wb.active
    ws.title = "Report"
    ws.sheet_view.showGridLines = False
    for i in range(1, ncols + 1):
        ws.column_dimensions[get_column_letter(i)].width = 16
    ws.column_dimensions["A"].width = 22

    # Header: company name in text, no logo (owner).
    ws.merge_cells(f"A1:{last}1")
    c = ws["A1"]; c.value = COMPANY.upper()
    style_cell(c, bold=True, size=16, color="FFFFFF", fill=HEAD_FILL, align=Alignment(horizontal="center", vertical="center"), border=False)
    ws.row_dimensions[1].height = 30
    ws.merge_cells(f"A2:{last}2")
    c = ws["A2"]; c.value = title
    style_cell(c, bold=True, size=13, align=Alignment(horizontal="center", vertical="center"), border=False)
    ws.row_dimensions[2].height = 22
    half = max(2, ncols // 2)
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=half)
    ws.cell(3, 1, "Report no.:").font = Font(name=F, size=10, bold=True)
    ws.merge_cells(start_row=3, start_column=half + 1, end_row=3, end_column=ncols)
    ws.cell(3, half + 1, "Page ____ of ____").font = Font(name=F, size=10, bold=True)
    ws.merge_cells(f"A4:{last}4")
    c = ws["A4"]; c.value = "Fill in the yellow cells. The 'Example' sheet shows the expected format. Check every field against your procedure, written practice and client specification before use."
    style_cell(c, size=9, color="555555", border=False)
    ws.row_dimensions[4].height = 26

    row = 6
    grid_info = None
    for kind, name, items in sections:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
        c = ws.cell(row, 1, name.upper())
        style_cell(c, bold=True, size=10, color="FFFFFF", fill=HEAD_FILL, align=Alignment(vertical="center"), border=False)
        row += 1
        if kind == "fields":
            split = max(2, ncols // 3)
            for label in items:
                ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=split)
                ws.merge_cells(start_row=row, start_column=split + 1, end_row=row, end_column=ncols)
                style_cell(ws.cell(row, 1, label), bold=True, size=9, fill=LABEL_FILL)
                for col in range(2, split + 1):
                    ws.cell(row, col).border = BORDER
                for col in range(split + 1, ncols + 1):
                    ws.cell(row, col).border = BORDER
                    ws.cell(row, col).fill = INPUT_FILL
                ws.cell(row, split + 1).alignment = WRAP
                ws.cell(row, split + 1).font = Font(name=F, size=10)
                ws.row_dimensions[row].height = 20
                row += 1
        elif kind == "box":
            ws.merge_cells(start_row=row, start_column=1, end_row=row + 3, end_column=ncols)
            for r in range(row, row + 4):
                for col in range(1, ncols + 1):
                    ws.cell(r, col).border = BORDER
                    ws.cell(r, col).fill = INPUT_FILL
            ws.cell(row, 1).alignment = WRAP
            row += 4
        else:  # grid
            cols = list(items)
            if slug == "ut-thickness-survey-report":
                cols = cols[:-1] + ["t-min (from design / FFS)", "Below t-min?", cols[-1]]
            for j, h in enumerate(cols, start=1):
                style_cell(ws.cell(row, j, h), bold=True, size=9, fill=LABEL_FILL, align=Alignment(wrap_text=True, horizontal="center", vertical="center"))
            ws.row_dimensions[row].height = 32
            hdr = row
            row += 1
            first = row
            for _ in range(18):
                for j in range(1, len(cols) + 1):
                    cell = ws.cell(row, j)
                    cell.border = BORDER
                    cell.fill = INPUT_FILL
                    cell.font = Font(name=F, size=10)
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
                if slug == "ut-thickness-survey-report":
                    # Minimum of the four directional readings; flag readings below t-min.
                    ws.cell(row, 8).value = f'=IF(COUNT(D{row}:G{row})=0,"",MIN(D{row}:G{row}))'
                    ws.cell(row, 8).fill = PatternFill(fill_type=None)
                    ws.cell(row, 10).value = f'=IF(OR(H{row}="",I{row}=""),"",IF(H{row}<I{row},"YES","no"))'
                    ws.cell(row, 10).fill = PatternFill(fill_type=None)
                    ws.cell(row, 8).number_format = "0.00"
                row += 1
            grid_info = (hdr, first, row - 1, cols)
            # Summary formulas under the table.
            if slug == "ut-thickness-survey-report":
                ws.cell(row, 1, "Lowest reading on this sheet (mm):").font = Font(name=F, size=9, bold=True)
                ws.cell(row, 8).value = f'=IF(COUNT(H{first}:H{row-1})=0,"",MIN(H{first}:H{row-1}))'
                ws.cell(row, 8).number_format = "0.00"
                ws.cell(row + 1, 1, "Locations below t-min:").font = Font(name=F, size=9, bold=True)
                ws.cell(row + 1, 8).value = f'=COUNTIF(J{first}:J{row-1},"YES")'
                row += 2
            else:
                res_col = next((j for j, h in enumerate(cols, start=1) if h.strip().lower() == "result"), None)
                ws.cell(row, 1, "Lines recorded:").font = Font(name=F, size=9, bold=True)
                ws.cell(row, 2).value = f"=COUNTA(A{first}:A{row-1})"
                if res_col:
                    L = get_column_letter(res_col)
                    ws.cell(row, 3, "Rejected:").font = Font(name=F, size=9, bold=True)
                    ws.cell(row, 4).value = f'=COUNTIF({L}{first}:{L}{row-1},"Reject")'
                    dv = DataValidation(type="list", formula1='"Accept,Reject,Repair,Re-examine"', allow_blank=True)
                    dv.add(f"{L}{first}:{L}{row-1}")
                    ws.add_data_validation(dv)
                row += 1
        row += 1

    # Footer
    ws.merge_cells(start_row=row, start_column=1, end_row=row + 1, end_column=ncols)
    c = ws.cell(row, 1, f"Template issued free by {COMPANY} ({SITE}). It is a starting point: your procedure, written practice and client specification decide what the report must contain. "
                        f"For templates in your own format, industry, code or method, request a quotation: {QUOTE_URL}")
    style_cell(c, size=8, color="555555", border=False)
    last_row = row + 1

    ws.print_area = f"A1:{last}{last_row}"
    ws.page_setup.orientation = "landscape" if ncols > 8 else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.oddFooter.center.text = f"{COMPANY} · {title} · Page &P of &N"

    # Example sheet
    ex = wb.create_sheet("Example")
    ex.sheet_view.showGridLines = False
    ex.column_dimensions["A"].width = 48
    ex.column_dimensions["B"].width = 60
    ex["A1"] = f"{title} — example values (format guide only; not a real job)"
    ex["A1"].font = Font(name=F, size=12, bold=True)
    r = 3
    for kind, name, items in sections:
        if kind == "fields":
            ex.cell(r, 1, name).font = Font(name=F, size=10, bold=True, color="1F3A5F"); r += 1
            for label in items:
                style_cell(ex.cell(r, 1, label), bold=True, size=9, fill=LABEL_FILL)
                cc = ex.cell(r, 2, ex_fields.get(label, "")); style_cell(cc, size=9)
                r += 1
            r += 1
    if grid_info and ex_grid:
        cols = grid_info[3]
        name = next(n for k, n, _ in sections if k == "grid")
        ex.cell(r, 1, name).font = Font(name=F, size=10, bold=True, color="1F3A5F"); r += 1
        src_cols = next(c for k, _, c in sections if k == "grid")
        for j, h in enumerate(src_cols, start=1):
            style_cell(ex.cell(r, j, h), bold=True, size=9, fill=LABEL_FILL)
            ex.column_dimensions[get_column_letter(j)].width = max(ex.column_dimensions[get_column_letter(j)].width or 0, 16)
        r += 1
        for vals in ex_grid:
            for j, v in enumerate(vals[:len(src_cols)], start=1):
                style_cell(ex.cell(r, j, v), size=9)
            r += 1

    # How to use / more templates
    hu = wb.create_sheet("Read me")
    hu.sheet_view.showGridLines = False
    hu.column_dimensions["A"].width = 110
    lines = [
        (f"{title}", Font(name=F, size=13, bold=True)),
        (f"Free template from {COMPANY} ({SITE}).", Font(name=F, size=10)),
        ("", None),
        ("How to use", Font(name=F, size=11, bold=True)),
        ("1. Put your company name, logo and document number in the header before first use.", Font(name=F, size=10)),
        ("2. Fill in the yellow cells. Formulas (where present) calculate the summary lines; do not overwrite them.", Font(name=F, size=10)),
        ("3. Check every field against your procedure, written practice and client specification, and add what they require.", Font(name=F, size=10)),
        ("4. The 'Example' sheet shows the expected format only; it is not data from a real job.", Font(name=F, size=10)),
        ("", None),
        ("Need more templates?", Font(name=F, size=11, bold=True)),
        ("These templates are a free starter set. We also prepare templates matched to your industry, method, code and client format —", Font(name=F, size=10)),
        ("for example PAUT/TOFD data sheets, IRIS and eddy-current tube reports, MFL tank-floor reports, hardness, PMI and holiday-test records,", Font(name=F, size=10)),
        ("calibration logs and complete report packs aligned to your written practice. Tell us your method, code and industry and we will send a quotation:", Font(name=F, size=10)),
        (QUOTE_URL, Font(name=F, size=10, color="1F5FBF", underline="single")),
    ]
    for i, (t, f) in enumerate(lines, start=1):
        hu.cell(i, 1, t)
        if f:
            hu.cell(i, 1).font = f
    hu.cell(len(lines), 1).hyperlink = QUOTE_URL

    wb.properties.creator = COMPANY
    wb.properties.title = title
    wb.properties.company = COMPANY if hasattr(wb.properties, "company") else None
    OUT.mkdir(parents=True, exist_ok=True)
    path = OUT / f"{slug}.xlsx"
    wb.save(path)
    return path


if __name__ == "__main__":
    for s in SLUGS:
        p = build(s)
        print("wrote", p.relative_to(ROOT))
